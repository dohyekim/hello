
import openpyxl
from PIL import Image
from openpyxl.chart import Reference, BarChart

book = openpyxl.load_workbook("./melon_top_xls.xlsx")
sheet = book.create_sheet()
sheet.title = "Chart"
sheet3 = book.worksheets[2]


rows = [
    ['김일수', 55],
    ['김이수', 11],
    ['김삼수', 33],
    ['김사수', 15],
    ['김오수', 11],
]

for row in rows:
    sheet.append(row)

# y축
data = Reference(sheet, min_col=2, min_row=1, max_col=2, max_row=5)

# x축    
categs = Reference(sheet, min_col=1, min_row=1, max_row=5)

chart = BarChart()
chart.add_data(data=data)
chart.set_categories(categs)

chart.legend = None  # 범례
chart.varyColors = True
chart.title = "Sample Chart"

sheet.add_chart(chart, "A1")


book.save("./melon_top_xls.xlsx")